import csv
import json

import pandas as pd
from docx import Document
from docx.shared import Pt

# Load the JSON data
with open('data-exclution-test-results/with-and-without-hybrid-int-clf-5-runs-all-configs/results.json') as f:
    data = json.load(f)

# Prepare the CSV data
csv_data = [['Configuration', '0% exc', '20% exc', '40% exc', '60% exc', '80% exc']]

# Temporary list to hold the configurations and their averages
temp_data = []

for config_name, values in data.items():
    # Calculate the average for each exclusion case and append a '%' sign
    averages = [str(round(sum(x[i] for x in values) / len(values) * 100, 2)) + '%' for i in range(5)]
    temp_data.append([config_name] + averages)

# Sort the configurations by the 0% exclusion case in descending order
sorted_data = sorted(temp_data, key=lambda x: float(x[1].strip('%')), reverse=True)

# Append the sorted data to the CSV data
csv_data.extend(sorted_data)

# Write the CSV data
with open('all_with_and_without_zero_shot.csv', 'w', newline='') as f:
    writer = csv.writer(f)
    writer.writerows(csv_data)

# Load the CSV data into a pandas DataFrame
df = pd.read_csv('all_with_and_without_zero_shot.csv')

# Create a new Word document
doc = Document()

# Add a table to the document
table = doc.add_table(df.shape[0] + 1, df.shape[1])

# Add the column names to the table
for j in range(df.shape[-1]):
    table.cell(0, j).text = df.columns[j]

# Add the rest of the data to the table
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        table.cell(i + 1, j).text = str(df.values[i, j])

# Save the document
doc.save('output_all_with_and_without_zero_shot.docx')
