import pandas as pd
from docx import Document
from docx.shared import Pt

# Load the CSV data into a pandas DataFrame
df = pd.read_csv('output_ablation_enh.csv')

# Create a new Word document
doc = Document()

# Add a table to the document
table = doc.add_table(df.shape[0]+1, df.shape[1])

# Add the column names to the table
for j in range(df.shape[-1]):
    table.cell(0,j).text = df.columns[j]

# Add the rest of the data to the table
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        table.cell(i+1,j).text = str(df.values[i,j])

# Save the document
doc.save('output.docx')
